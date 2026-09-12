"""使用临时文件验证 PDF、DOCX 与 GIS ZIP 提取，不写入正式知识库。"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfWriter


EXTRACTOR = Path(__file__).with_name("extract_source.py")


def extract(path: Path) -> dict:
    result = subprocess.run(
        [sys.executable, str(EXTRACTOR), str(path)],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    return json.loads(result.stdout)


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="nwu-extract-test-") as folder:
        root = Path(folder)

        pdf_path = root / "记录.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        with pdf_path.open("wb") as stream:
            writer.write(stream)
        assert extract(pdf_path)["metadata"]["pages"] == 1

        docx_path = root / "记录.docx"
        document = Document()
        document.add_heading("青岚岭记录", level=1)
        document.add_paragraph("北坡观察点可见砂岩层。")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "经度"
        table.rows[0].cells[1].text = "116.123"
        document.save(docx_path)
        docx_result = extract(docx_path)
        assert "北坡观察点" in docx_result["text"]
        assert docx_result["metadata"]["tables"] == 1

        geojson_path = root / "points.geojson"
        geojson_path.write_text(json.dumps({
            "type": "FeatureCollection",
            "features": [{
                "type": "Feature",
                "geometry": {"type": "Point", "coordinates": [116.123, 29.456]},
                "properties": {"name": "青岚岭北坡观察点"},
            }],
        }, ensure_ascii=False), encoding="utf-8")
        zip_path = root / "gis.zip"
        with zipfile.ZipFile(zip_path, "w") as archive:
            archive.write(geojson_path, geojson_path.name)
        gis_result = extract(zip_path)
        assert gis_result["metadata"]["layers"] == 1
        assert "青岚岭北坡观察点" in gis_result["text"]

    print("资料提取测试通过：PDF、DOCX、GIS ZIP 均可读取")


if __name__ == "__main__":
    main()
