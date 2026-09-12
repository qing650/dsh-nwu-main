#!/usr/bin/env python3
"""西北大学教务问答资料提取器。stdout 只输出一个 UTF-8 JSON 对象。"""

from __future__ import annotations

import json
import math
import os
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable


MAX_ARCHIVE_FILES = 20_000
MAX_ARCHIVE_BYTES = 2 * 1024 * 1024 * 1024
MAX_GIS_FEATURES = 20_000


def fail(message: str) -> None:
    raise RuntimeError(message)


def pdf(path: Path) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
    except ImportError:
        fail("缺少 pypdf，请先运行“安装教务问答.bat”")
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages):
        pages.append(f"\n\n## 第 {index + 1} 页\n{page.extract_text() or ''}")
    return {"text": "".join(pages).strip(), "metadata": {"pages": len(reader.pages)}}


def docx(path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError:
        fail("缺少 python-docx，请先运行“安装教务问答.bat”")
    document = Document(str(path))
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text.strip())
    for table_index, table in enumerate(document.tables, 1):
        chunks.append(f"\n表 {table_index}")
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return {
        "text": "\n".join(chunks),
        "metadata": {"paragraphs": len(document.paragraphs), "tables": len(document.tables)},
    }


def audio(path: Path) -> dict[str, Any]:
    try:
        from faster_whisper import WhisperModel
    except ImportError:
        fail("缺少 faster-whisper，请先运行“安装教务问答.bat”")
    model_name = os.environ.get("NWU_WHISPER_MODEL", "small")
    device = os.environ.get("NWU_WHISPER_DEVICE", "cpu")
    compute_type = os.environ.get("NWU_WHISPER_COMPUTE", "int8" if device == "cpu" else "float16")
    model = WhisperModel(model_name, device=device, compute_type=compute_type)
    segments, info = model.transcribe(str(path), language="zh", vad_filter=True, beam_size=5)
    rows = []
    for segment in segments:
        text = segment.text.strip()
        if text:
            rows.append(f"[{segment.start:0.1f}—{segment.end:0.1f} 秒] {text}")
    return {
        "text": "\n".join(rows),
        "metadata": {
            "language": info.language,
            "languageProbability": info.language_probability,
            "duration": info.duration,
            "model": model_name,
        },
    }


def safe_extract(archive: Path, target: Path) -> None:
    with zipfile.ZipFile(archive) as package:
        members = package.infolist()
        if len(members) > MAX_ARCHIVE_FILES:
            fail(f"ZIP 文件数超过 {MAX_ARCHIVE_FILES} 限制")
        if sum(member.file_size for member in members) > MAX_ARCHIVE_BYTES:
            fail("ZIP 解压后超过 2 GB 限制")
        root = target.resolve()
        for member in members:
            destination = (target / member.filename).resolve()
            if destination != root and root not in destination.parents:
                fail("ZIP 包含越界路径，已拒绝解压")
        package.extractall(target)


def iter_positions(value: Any) -> Iterable[tuple[float, float]]:
    if isinstance(value, (list, tuple)):
        if len(value) >= 2 and isinstance(value[0], (int, float)) and isinstance(value[1], (int, float)):
            yield float(value[0]), float(value[1])
        else:
            for item in value:
                yield from iter_positions(item)


def geometry_summary(geometry: Any) -> tuple[list[float] | None, list[float] | None]:
    if not geometry:
        return None, None
    points = []
    for point in iter_positions(geometry.get("coordinates")):
        points.append(point)
        if len(points) >= 100_000:
            break
    if not points:
        return None, None
    xs = [point[0] for point in points]
    ys = [point[1] for point in points]
    bbox = [min(xs), min(ys), max(xs), max(ys)]
    center = [(bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2]
    if not all(math.isfinite(value) for value in bbox + center):
        return None, None
    return bbox, center


def layer_sources(root: Path) -> list[tuple[str, str | None]]:
    sources: list[tuple[str, str | None]] = []
    for gdb in root.rglob("*.gdb"):
        try:
            import fiona
            for layer in fiona.listlayers(str(gdb)):
                sources.append((str(gdb), layer))
        except Exception as error:
            fail(f"无法读取 GDB {gdb.name}：{error}")
    for shp in root.rglob("*.shp"):
        sources.append((str(shp), None))
    for geojson in [*root.rglob("*.geojson"), *root.rglob("*.json")]:
        sources.append((str(geojson), None))
    return sources


def gis_zip(path: Path) -> dict[str, Any]:
    try:
        import fiona
    except ImportError:
        fail("缺少 Fiona，请先运行“安装教务问答.bat”")
    temp = Path(tempfile.mkdtemp(prefix="nwu-gis-"))
    try:
        safe_extract(path, temp)
        layers = []
        lines = []
        total = 0
        for source_path, layer_name in layer_sources(temp):
            kwargs = {"layer": layer_name} if layer_name else {}
            with fiona.open(source_path, **kwargs) as collection:
                label = layer_name or Path(source_path).stem
                count = len(collection)
                schema = dict(collection.schema)
                record = {
                    "name": label,
                    "featureCount": count,
                    "crs": collection.crs.to_string() if collection.crs else None,
                    "schema": schema,
                    "features": [],
                }
                lines.append(f"\n## 图层 {label}\n要素数：{count}\n坐标系：{record['crs'] or '未声明'}\n字段：{json.dumps(schema.get('properties', {}), ensure_ascii=False)}")
                for index, feature in enumerate(collection):
                    if total >= MAX_GIS_FEATURES:
                        break
                    properties = dict(feature.get("properties") or {})
                    geometry = feature.get("geometry")
                    bbox, center = geometry_summary(geometry)
                    title = properties.get("name") or properties.get("NAME") or properties.get("Name") or properties.get("名称") or f"{label} 要素 {index + 1}"
                    item = {
                        "title": str(title),
                        "geometryType": geometry.get("type") if geometry else None,
                        "center": center,
                        "bbox": bbox,
                        "properties": properties,
                    }
                    record["features"].append(item)
                    lines.append(f"{index + 1}. {item['title']}｜{item['geometryType'] or '无几何'}｜中心 {center or '无'}｜属性 {json.dumps(properties, ensure_ascii=False, default=str)}")
                    total += 1
                layers.append(record)
        if not layers:
            fail("ZIP 中没有找到可读取的 GDB、Shapefile 或 GeoJSON")
        if total >= MAX_GIS_FEATURES:
            lines.append(f"\n仅提取前 {MAX_GIS_FEATURES} 个要素，原文件仍已完整保存。")
        return {
            "text": "\n".join(lines).strip(),
            "structured": {"format": "GIS archive", "layers": layers, "extractedFeatures": total},
            "metadata": {"layers": len(layers), "extractedFeatures": total},
        }
    finally:
        shutil.rmtree(temp, ignore_errors=True)


def main() -> None:
    if len(sys.argv) != 2:
        fail("用法：extract_source.py <文件路径>")
    path = Path(sys.argv[1]).resolve()
    if not path.is_file():
        fail("输入文件不存在")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        result = pdf(path)
    elif suffix == ".docx":
        result = docx(path)
    elif suffix in {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac"}:
        result = audio(path)
    elif suffix == ".zip":
        result = gis_zip(path)
    else:
        fail(f"提取器不支持 {suffix} 文件")
    print(json.dumps(result, ensure_ascii=False, default=str))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        raise SystemExit(1)
